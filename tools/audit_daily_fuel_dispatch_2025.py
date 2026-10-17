"""Independent daily-source, resource cost, native physics and full-year replay."""
import argparse,csv,datetime as dt,json,logging,resource,signal,time
from pathlib import Path
from zoneinfo import ZoneInfo
import numpy as np
import xlrd
from openpyxl import load_workbook
from generation_mix_2025 import ROOT,MONTH,CARRIER,EMBER,digest
from bidding_zone_network_2025 import build
from global_bidding_zone_2025 import compile_global
from nve_hydro_diagnostics import branch_offsets
from bid_calibration_2025 import immutable_hashes
from audit_nve_hydro_diagnostics_2025 import physics
from price_mechanism_2025 import dual_replay
from audit_seasonal_calendar_bids_2025 import compare,price_summary


def source_replay(folder,workbook,law):
    s=json.loads((folder/'summary.json').read_text())
    assert s['producer_sha256']==digest(ROOT/'tools/daily_fuel_inputs_2025.py')
    assert s['eia']['workbook_sha256']==digest(workbook) and s['inputs_sha256']==digest(folder/'inputs.npz')
    for p,h in s['sources'].items():assert digest(ROOT/p)==h
    wb=xlrd.open_workbook(str(workbook));sheet=wb.sheet_by_name('Data 1');quotes={}
    for i in range(3,sheet.nrows):
        if not sheet.cell_value(i,0):continue
        day=xlrd.xldate_as_datetime(sheet.cell_value(i,0),wb.datemode).date();v=sheet.cell_value(i,1)
        if day.year==2025 and isinstance(v,(float,int)):
            assert day not in quotes;quotes[day]=v
    with (folder/'brent-quotes-2025.csv').open() as f:cached=list(csv.DictReader(f))
    assert s['eia']['quote_csv_sha256']==digest(folder/'brent-quotes-2025.csv')
    assert len(cached)==len(quotes)==s['eia']['quote_count']
    for row in cached:assert quotes[dt.date.fromisoformat(row['date'])]==float(row['brent_usd_barrel'])
    gasroot=ROOT/'data/fuel-prices/mibgas-daily-source-v1';g=json.loads((gasroot/'receipt.json').read_text())
    assert g['raw_workbook_sha256']==digest(gasroot/'source-2025.xlsx')
    ws=load_workbook(gasroot/'source-2025.xlsx',read_only=True,data_only=True)['MIBGAS Indexes'];cells=ws.iter_rows(values_only=True);header=next(cells)
    banks={key:{} for key in ['ES','PT']}
    for row in cells:
        if row[0] is None:continue
        day=row[0].date()
        for zone,hub in [('ES','PVB'),('PT','VTP')]:banks[zone][day]=float(row[header.index(f'MIBGAS {hub} Last Price Index Day-Ahead\n[EUR/MWh]')])
    old=json.loads((ROOT/'data/daily-market-2025/monthly-fuel-annual-v2/fuel-inputs.json').read_text())
    fx={}
    with (ROOT/'data/fuel-prices/ecb-2025-api.csv').open() as f:
        for row in csv.DictReader(f):assert row['KEY']=='EXR.D.USD.EUR.SP00.A';fx[dt.date.fromisoformat(row['TIME_PERIOD'])]=float(row['OBS_VALUE'])
    with np.load(folder/'inputs.npz') as z:inputs={k:z[k] for k in z.files}
    # Construct delivery intervals independently; this explicitly covers 23/25-hour gas days.
    start=dt.datetime(2025,1,1,tzinfo=dt.timezone.utc);gasdates=['']*8760;durations=[]
    for day in sorted(banks['ES']):
        a=dt.datetime.combine(day,dt.time(6),ZoneInfo('Europe/Madrid')).astimezone(dt.timezone.utc)
        b=dt.datetime.combine(day+dt.timedelta(days=1),dt.time(6),ZoneInfo('Europe/Madrid')).astimezone(dt.timezone.utc)
        left=int((a-start).total_seconds()/3600);right=int((b-start).total_seconds()/3600);durations.append(right-left)
        for h in range(max(left,0),min(right,8760)):gasdates[h]=day.isoformat()
    assert 23 in durations and 25 in durations
    np.testing.assert_array_equal(gasdates,inputs['gas_delivery_day'])
    for t in range(8760):
        date=(start+dt.timedelta(hours=t)).date();qd=max((d for d in quotes if d<=date),default=None);fd=max((d for d in fx if d<=date),default=None)
        expected=old['hours'][t]['oil_eur_mwh_th']
        if qd is not None and fd is not None:
            assert (date-qd).days<=5 and (date-fd).days<=5
            expected=quotes[qd]/fx[fd]/old['assumptions']['oil_mwh_th_per_barrel']
            assert inputs['oil_quote_date'][t]==qd.isoformat() and inputs['fx_quote_date'][t]==fd.isoformat()
        else:assert inputs['oil_quote_date'][t]==inputs['fx_quote_date'][t]==''
        assert abs(expected-inputs['oil'][t])<1e-10
        for zone in banks:
            expected=banks[zone][dt.date.fromisoformat(gasdates[t])]*old['assumptions']['gas_basis_multiplier'] if gasdates[t] else old['hours'][t]['gas_eur_mwh_th']
            assert abs(expected-inputs['gas_'+zone][t])<1e-10
    assert np.count_nonzero(inputs['oil_quote_date']=='')==s['oil_monthly_fallback_hours']==24
    assert gasdates.count('')==s['gas_monthly_fallback_hours_per_zone']==5
    return inputs,old,dict(status='daily_brent_fx_and_original_gas_index_cells_and_dst_replayed',gas_day_law_url='https://eur-lex.europa.eu/legal-content/EN/TXT/PDF/?uri=CELEX:32017R0459',gas_day_law_sha256=digest(law),oil_daily_hours=8736,gas_daily_hours_per_zone=8755)


def audit(folder,workbook,law,out):
    began=time.perf_counter();s=json.loads((folder/'summary.json').read_text())
    assert s['producer_sha256']==digest(ROOT/'tools/run_daily_fuel_dispatch_2025.py')
    for name,h in s['source_hashes'].items():assert digest(ROOT/name)==h
    for name,h in s['dependencies'].items():assert digest(ROOT/'tools'/name)==h
    compare(s['design'],json.loads((folder/'design.json').read_text()))
    daily,old,source=source_replay(ROOT/s['inputs_folder'],workbook,law)
    root=ROOT/s['accepted_folder'];accepted=json.loads((root/'summary.json').read_text())
    n,raw,fixed,*_=build();base,off=compile_global(n,raw,fixed,'bidding_zone');flow=branch_offsets(n,off)
    assert immutable_hashes(base,flow)==s['original_physical_inputs']
    layout=json.loads((root/'layout.json').read_text());parent=np.array(layout['parent']);fractions=np.array(layout['fractions']);members=layout['members']
    available=n.get_switchable_as_dense('Generator','p_max_pu').values*n.generators.p_nom.values
    m=base.copy();m['ng']=len(parent);m['A']=base['A'][:,np.r_[parent,np.arange(base['ng'],base['A'].shape[1])]].tocsc()
    m['availability']=np.column_stack([available[:,ids].sum(axis=1)*f for ids,f in zip(members,fractions)]);m['cost']=base['cost'][:,parent]
    for key in ['offer_zones','emergency_mask']:m[key]=base[key][parent]
    assert immutable_hashes(m,flow)==s['expanded_physical_inputs']
    proposal=json.loads((root/'proposal-refit.json').read_text());beta=np.array(proposal['beta']);gamma=np.array(proposal['gamma'])
    with np.load(root/'calendar.npz') as z:features=z['features'];season=z['season'];shares=z['country_shares']
    with np.load(root/'development-targets.npz') as z:target={k:z[k] for k in z.files}
    with np.load(root/'evaluation-targets.npz') as z:target={k:np.concatenate([v,z[k]]) for k,v in target.items()}
    countries=s['countries'];fuels=['Wind','Solar','Nuclear','Gas','Coal'];fmap={v:k for k,vs in EMBER.items() for v in vs};labels=[(n.buses.loc[g.bus,'country'],fmap.get(CARRIER[g.carrier])) for _,g in n.generators.iterrows()]
    water=np.zeros((8760,len(countries)));demand=np.zeros_like(water);loads=n.get_switchable_as_dense('Load','p_set')
    for ci,country in enumerate(countries):
        buses=n.buses.index[n.buses.country==country];water[:,ci]=fixed[buses].sum(axis=1);demand[:,ci]=loads[n.loads.index[n.loads.bus.isin(buses)]].sum(axis=1)
    gas=np.array([r['gas_eur_mwh_th'] for r in old['hours']]);oil=np.array([r['oil_eur_mwh_th'] for r in old['hours']]);certs=[]
    for case,info in s['annuals'].items():
        p=ROOT/info['witness'] if info.get('external') else folder/info['witness'];assert digest(p)==info['witness_sha256']
        with np.load(p) as z:x=z['values'];y=z['row_duals'];obj=z['objectives'];assert z['optimal'].all()
        cost=m['cost'].copy();curves={k:daily[k].copy() for k in ['oil','gas_ES','gas_PT']}
        if case=='both_daily_shape_only':
            for key in curves:
                reference=oil if key=='oil' else gas
                for month in range(1,13):
                    ids=np.flatnonzero(MONTH==month);curves[key][ids]=curves[key][ids]/np.mean(curves[key][ids])*reference[ids[0]]
        for j,ids in enumerate(members):
            values=[]
            for i in ids:
                g=n.generators.iloc[i];v=raw[:,i].copy();zone=base['buszone'][g.bus]
                if case!='baseline':
                    if g.carrier=='oil' and case!='gas_daily_ES_PT':v+=(curves['oil']-oil)/g.efficiency
                    elif g.carrier in ['CCGT','OCGT'] and zone in ['ES','PT'] and case!='oil_daily':v+=(curves['gas_'+zone]-gas)/g.efficiency
                values.append(v)
            assert all(np.allclose(v,values[0],atol=1e-8,rtol=1e-12) for v in values)
            cost[:,j]=values[0];family=layout['kinds'][j]
            if family in ['gas','coal','nuclear']:
                zi=m['zones'].index(m['offer_zones'][j]);fi=['gas','coal','nuclear'].index(family)
                correction=(features[:,zi]*beta[zi]).sum(axis=1)+(season@gamma[:,fi].T)@shares[j]
                cost[:,j]=np.maximum(cost[:,j]+np.clip(correction,-30,30)*(.25 if family=='nuclear' else 1),0)
        original=np.zeros((8760,base['ng']))
        for j,k in enumerate(parent):original[:,k]+=x[:,j]
        original=np.c_[original,x[:,m['ng']:]];residual=physics(n,base,original,base['load'],off.copy(),0)
        error=float(abs((x[:,:m['ng']]*cost).sum(axis=1)-obj).max());assert error<.05
        reduced=np.c_[cost,np.zeros((8760,m['nl']+m['nz']))]-(m['A'].T@y.T).T
        dual=dual_replay(m|dict(cost=cost),x,y,reduced,obj,flow)
        compare(price_summary(y[:,:m['nz']],target['prices'],m['zones']),s['outcomes'][case]['prices'])
        lo=np.zeros((12,len(countries),5));hi=np.zeros_like(lo);cl=np.zeros((12,len(countries)));cu=np.zeros_like(cl);cp=np.zeros_like(cl)
        def add(v):return np.bincount(MONTH,weights=v,minlength=13)[1:]/1e6
        for j,(ids,f) in enumerate(zip(members,fractions)):
            a=available[:,ids].sum(axis=1)*f;q=np.clip(x[:,j],0,a)
            for c in {labels[i][0] for i in ids}:
                ci=countries.index(c);real=[i for i in ids if labels[i][0]==c and n.generators.iloc[i].carrier!='emergency']
                cap=available[:,real].sum(axis=1)*f
                cl[:,ci]+=add(np.maximum(q-a+cap,0));cu[:,ci]+=add(np.minimum(q,cap));cp[:,ci]+=add(np.divide(q*cap,a,out=np.zeros(8760),where=a>0))
                for fuel in fuels:
                    selected=[i for i in real if labels[i][1]==fuel]
                    if not selected:continue
                    cap=available[:,selected].sum(axis=1)*f;fi=fuels.index(fuel)
                    lo[:,ci,fi]+=add(np.maximum(q-a+cap,0));hi[:,ci,fi]+=add(np.minimum(q,cap))
        dm=np.column_stack([add(demand[:,i]) for i in range(len(countries))]);w=np.column_stack([add(water[:,i]) for i in range(len(countries))]);cl+=w-dm;cu+=w-dm;cp+=w
        monthly=[dict(month=i+1,generation_lower_twh=lo[i].tolist(),generation_upper_twh=hi[i].tolist(),generation_point_twh=cp[i].tolist(),net_export_lower_twh=cl[i].tolist(),net_export_upper_twh=cu[i].tolist(),net_export_point_twh=(cp-dm)[i].tolist(),demand_twh=dm[i].tolist()) for i in range(12)]
        compare(monthly,s['outcomes'][case]['monthly_quantities'])
        ref=target['generation'];tr=target['trade'];known=np.isfinite(ref);tk=np.isfinite(tr)
        quantity=dict(price=s['outcomes'][case]['prices']['equal_zone_mae'],generation=float((np.maximum(np.maximum(lo-ref,ref-hi),0)/np.maximum(ref,.25))[known].mean()),trade=float((np.maximum(np.maximum(cl-tr,tr-cu),0)/np.maximum(dm,1))[tk].mean()),price_areas=41,generation_targets=int(known.sum()),trade_targets=int(tk.sum()))
        compare(quantity,s['outcomes'][case]['quantity_errors'])
        emergency=float(x[:,:m['ng']][:,m['emergency_mask']].sum());assert abs(emergency-info['emergency_mwh'])<1e-8
        oil_lo=oil_hi=0.
        for j,(ids,f) in enumerate(zip(members,fractions)):
            selected=[i for i in ids if n.generators.iloc[i].carrier=='oil']
            if not selected:continue
            total=available[:,ids].sum(axis=1)*f;cap=available[:,selected].sum(axis=1)*f;q=np.clip(x[:,j],0,total)
            oil_lo+=float(np.maximum(q-total+cap,0).sum()/1e6);oil_hi+=float(np.minimum(q,cap).sum()/1e6)
        certs.append(dict(case=case,hours=8760,original_physical_residual_mw=residual,objective_error_eur=error,oil_generation_lower_twh=oil_lo,oil_generation_upper_twh=oil_hi,**dual));print('Replayed',case,flush=True)
    assert immutable_hashes(m,flow)==s['expanded_physical_inputs']
    out.mkdir(parents=True,exist_ok=False);result=dict(status='daily_fuel_sources_costs_original_physics_duals_prices_and_monthly_mix_replayed',summary_sha256=digest(folder/'summary.json'),auditor_sha256=digest(__file__),source_replay=source,certificates=certs,elapsed_seconds=time.perf_counter()-began)
    (out/'audit.json').write_text(json.dumps(result,indent=2)+'\n');print(result['status'],flush=True)


if __name__=='__main__':
    logging.getLogger('pypsa').setLevel(logging.ERROR);p=argparse.ArgumentParser();p.add_argument('folder',type=Path);p.add_argument('workbook',type=Path);p.add_argument('law',type=Path);p.add_argument('--out',type=Path);a=p.parse_args();resource.setrlimit(resource.RLIMIT_AS,(5*1024**3,5*1024**3));signal.alarm(1800);audit(a.folder,a.workbook,a.law,a.out or a.folder/'verification-v1')
