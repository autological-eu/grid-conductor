"""Reparse original MIBGAS cells independently of the pandas extraction."""
import argparse,csv,datetime as dt,json,math
from pathlib import Path
from openpyxl import load_workbook
from generation_mix_2025 import ROOT,digest


def audit(raw,folder):
    r=json.loads((folder/'receipt.json').read_text())
    assert r['raw_workbook_sha256']==digest(raw)
    assert r['producer_sha256']==digest(ROOT/'tools/prepare_mibgas_daily_2025.py')
    assert r['output_sha256']==digest(folder/'daily-indices-2025.csv')
    book=load_workbook(raw,read_only=True,data_only=True)
    cells=book['MIBGAS Indexes'].iter_rows(values_only=True);header=next(cells);original={}
    for row in cells:
        if row[0] is None:continue
        day=row[0].date().isoformat();assert day not in original
        original[day]={name:row[header.index(column)] for name,column in r['columns'].items()}
    with (folder/'daily-indices-2025.csv').open() as f:rows=list(csv.DictReader(f))
    expected=[(dt.date(2025,1,1)+dt.timedelta(days=n)).isoformat() for n in range(365)]
    assert [v['delivery_day'] for v in rows]==expected and set(original)==set(expected)
    known={name:0 for name in r['columns']}
    for row in rows:
        for name in known:
            v=original[row['delivery_day']][name]
            if v is None:assert row[name]==''
            else:
                assert math.isfinite(v) and v>=0 and abs(float(row[name])-v)<1e-10
                known[name]+=1
    assert r['delivery_days']==365 and known==r['known_days']
    result=dict(status='original_365_gas_day_cells_replayed_not_adopted_in_model',receipt_sha256=digest(folder/'receipt.json'),auditor_sha256=digest(__file__),known_days=known)
    out=folder/'audit.json'
    if out.exists():raise FileExistsError(out)
    out.write_text(json.dumps(result,indent=2)+'\n');print(json.dumps(result,indent=2))


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('workbook',type=Path);p.add_argument('folder',type=Path);a=p.parse_args();audit(a.workbook,a.folder)
